"""Genera documentos sintéticos para verificar cada conversor incluido."""

from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook
from pptx import Presentation
from PySide6.QtGui import QFont, QPainter, QPdfWriter

from markitdown_es.__main__ import create_app

folder = Path(__file__).resolve().parents[1] / "tests" / "fixtures"
folder.mkdir(parents=True, exist_ok=True)
(folder / "ejemplo.txt").write_text(
    "Proyecto Aurora\nInformación para el equipo.", encoding="utf-8"
)
(folder / "ejemplo.html").write_text(
    "<html><body><h1>Proyecto Aurora</h1><p>Información para el equipo.</p><ul><li>Revisar documentos</li><li>Compartir resultados</li></ul></body></html>",
    encoding="utf-8",
)
(folder / "ejemplo.csv").write_text("Proyecto,Estado\nProyecto Aurora,En curso\n", encoding="utf-8")
with ZipFile(folder / "ejemplo.docx", "w") as archive:
    archive.writestr(
        "[Content_Types].xml",
        '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
    )
    archive.writestr(
        "_rels/.rels",
        '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
    )
    archive.writestr(
        "word/document.xml",
        '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Proyecto Aurora</w:t></w:r></w:p></w:body></w:document>',
    )
workbook = Workbook()
workbook.active.append(["Proyecto", "Estado"])
workbook.active.append(["Proyecto Aurora", "En curso"])
workbook.save(folder / "ejemplo.xlsx")
slides = Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[0])
slide.shapes.title.text = "Proyecto Aurora"
slides.save(folder / "ejemplo.pptx")
app = create_app()
writer = QPdfWriter(str(folder / "ejemplo.pdf"))
painter = QPainter(writer)
painter.setFont(QFont("Arial", 16))
painter.drawText(500, 500, "Proyecto Aurora")
painter.end()
print(f"7 documentos creados en {folder}")
